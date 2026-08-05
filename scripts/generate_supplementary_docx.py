"""
Generate submission-ready Supplementary Materials DOCX for
Parankusham et al. (2026) — GRAD: Journal of Prevention of Alzheimer's Disease
"""

import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os
from _grad_paths import (RESULTS, ADNI_DIR, A4_DIR, DATA_DIR, PROJECT_ROOT, SYNTHETIC, load_threshold_sweep)  # noqa: F401

OUTPUT = str(RESULTS / 'GRAD_Supplementary_Materials.docx')
RESULTS = str(RESULTS)
FIGURES = f'{RESULTS}/figures'
TABLES = f'{RESULTS}/tables'


def set_cell_shading(cell, color_hex):
    """Set cell background color."""
    shading = cell._element.get_or_add_tcPr()
    shading_elem = shading.makeelement(qn('w:shd'), {
        qn('w:fill'): color_hex,
        qn('w:val'): 'clear',
    })
    shading.append(shading_elem)


def add_styled_table(doc, df, caption):
    """Add a formatted table with header styling."""
    # Caption
    p = doc.add_paragraph()
    run = p.add_run(caption.split('.')[0] + '.')
    run.bold = True
    run.font.size = Pt(11)
    if '.' in caption:
        rest = '.'.join(caption.split('.')[1:])
        run2 = p.add_run(rest)
        run2.font.size = Pt(10)
    p.space_after = Pt(6)

    # Table
    table = doc.add_table(rows=1 + len(df), cols=len(df.columns))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Header row
    for j, col_name in enumerate(df.columns):
        cell = table.rows[0].cells[j]
        cell.text = str(col_name)
        set_cell_shading(cell, '2166AC')
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9)

    # Data rows
    for i, row_data in enumerate(df.values):
        for j, val in enumerate(row_data):
            cell = table.rows[i + 1].cells[j]
            cell.text = str(val)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(9)
            # Alternating row shading
            if i % 2 == 1:
                set_cell_shading(cell, 'F5F5F5')

    doc.add_paragraph()  # Spacer


def main():
    os.chdir(os.path.dirname(os.path.abspath(__file__)))

    doc = Document()

    # ── Page margins ─────────────────────────────────────────────────
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)

    # ── Title page ───────────────────────────────────────────────────
    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Additional File 1: Supplementary Materials')
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        'GRAD: A Two-Stage Algorithm for Resolving Diagnostic\n'
        'Uncertainty in the Plasma phospho-tau217 Gray Zone')
    run.font.size = Pt(14)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        'Harthik Parankusham, Casey Vanderlip, Colin Birkenbihl, Eashwar Krishna,\n'
        'Chizobam Ugboaja, Andrew Budson, Brandon Frank,\n'
        'and for the Alzheimer\'s Disease Neuroimaging Initiative')
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(68, 68, 68)

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Journal of Prevention of Alzheimer\'s Disease')
    run.font.size = Pt(11)
    run.italic = True
    run.font.color.rgb = RGBColor(102, 102, 102)

    doc.add_paragraph()
    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('Contents:\n\nSupplementary Figures S1\u2013S2\nSupplementary Tables S1\u2013S7')
    run.font.size = Pt(11)

    doc.add_page_break()

    # ── Figure S1: Calibration ───────────────────────────────────────
    p = doc.add_paragraph()
    run = p.add_run('Figure S1.')
    run.bold = True
    run.font.size = Pt(11)
    run = p.add_run(
        ' ADNI LOOCV calibration analysis. '
        '(A) Calibration curve comparing GRAD predicted probabilities '
        'against observed A\u03b2 positivity fraction across decile bins. '
        'Dashed line indicates perfect calibration. Expected Calibration '
        'Error (ECE) = 0.060. Hosmer\u2013Lemeshow test: \u03c7\u00b2 = 17.69, '
        'df = 8, p = 0.024. '
        '(B) Prediction density distributions by true A\u03b2 status, with '
        'Gatekeeper probability thresholds (P = 0.25, P = 0.75) indicated '
        'by dashed vertical lines.')
    run.font.size = Pt(10)
    p.space_after = Pt(12)

    doc.add_picture(f'{FIGURES}/figure_s1_calibration.png', width=Inches(6.0))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # ── Figure S2: Subgroup & Threshold (was S3, now S2) ─────────────
    p = doc.add_paragraph()
    run = p.add_run('Figure S2.')
    run.bold = True
    run.font.size = Pt(11)
    run = p.add_run(
        ' Supplementary model characterization. '
        '(A) Subgroup AUC forest plot stratified by cognitive status, '
        'APOE \u03b54 status, sex, and age tertiles; dashed line indicates '
        'overall AUC (0.857). '
        '(B) Threshold\u2013performance tradeoff showing sensitivity, '
        'specificity, PPV, NPV, and accuracy across classification '
        'thresholds, with 90% sensitivity (threshold = 0.240) and 90% '
        'specificity (threshold = 0.674) operating points annotated. '
        '(C) Gatekeeper threshold sensitivity heatmap across 25 '
        'low/high threshold combinations. Each cell shows resolution '
        'rate (top) and resolved accuracy (bottom). Gold border '
        'indicates the selected 0.25/0.75 operating point. '
        '(D) Per-stage confusion matrices for Stage 1 Gatekeeper '
        '(N = 178, accuracy 88.8%) and Stage 2 Reflex '
        '(N = 142, accuracy 70.4%).')
    run.font.size = Pt(10)
    p.space_after = Pt(12)

    doc.add_picture(f'{FIGURES}/figure_s3_subgroup_threshold.png',
                    width=Inches(6.5))
    last_paragraph = doc.paragraphs[-1]
    last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # ── Table S1: Complete Model Performance Metrics ─────────────────
    # Expanded to include all 10 metrics from the former main-text Table 2
    # (relocated here to satisfy JPAD's 5-graphic main-text limit).
    df_display = pd.DataFrame({
        'Metric': ['AUC', 'AUPRC', 'Accuracy', 'Sensitivity', 'Specificity',
                   'PPV', 'NPV', 'LR+', 'LR−', 'Brier Score'],
        'Value': ['0.857', '0.827', '80.6%', '78.7%', '82.4%',
                  '80.8%', '80.5%', '4.48', '0.26', '0.148'],
        '95% CI': ['0.813–0.897', '0.757–0.897',
                   '76.2%–84.7%', '72.3%–84.8%',
                   '76.4%–88.0%', '74.5%–86.7%',
                   '74.6%–86.2%', '3.28–6.50',
                   '0.18–0.35', '—'],
    })
    add_styled_table(
        doc, df_display,
        'Table S1. Complete Model Performance Metrics. '
        'Full GRAD model performance in the ADNI development cohort '
        '(N = 320) under leave-one-out cross-validation. 95% confidence '
        'intervals for AUC, accuracy, sensitivity, and specificity derived '
        'from 2,000 bootstrap resamples of the LOOCV predictions; CIs for '
        'PPV, NPV, and likelihood ratios computed analytically. '
        'Discrimination, calibration, and likelihood ratio metrics '
        'reported per STARD 2015 diagnostic accuracy guidelines. '
        'AUC, area under the receiver operating characteristic curve; '
        'AUPRC, area under the precision-recall curve; PPV, positive '
        'predictive value; NPV, negative predictive value; LR+, positive '
        'likelihood ratio; LR−, negative likelihood ratio.'
    )

    doc.add_page_break()

    # ── Table S2: Subgroup Performance ───────────────────────────────
    df = pd.read_csv(f'{TABLES}/supp_table_subgroup_performance.csv')
    df = df[df['subgroup'] != 'Overall']
    df['subgroup'] = df['subgroup'].replace({
        'Young tertile (\u226470 / \u226475 / >75)': 'Age \u226470',
        'Middle tertile (\u226470 / \u226475 / >75)': 'Age 70\u201375',
        'Old tertile (\u226470 / \u226475 / >75)': 'Age >75',
    })
    df_display = df[['category', 'subgroup', 'n', 'n_positive',
                     'auc', 'accuracy', 'sensitivity', 'specificity']].copy()
    df_display.columns = ['Category', 'Subgroup', 'N', 'N+',
                          'AUC', 'Accuracy', 'Sensitivity', 'Specificity']
    for c in ['AUC', 'Accuracy', 'Sensitivity', 'Specificity']:
        df_display[c] = df_display[c].apply(lambda x: f'{x:.3f}')
    add_styled_table(
        doc, df_display,
        'Table S2. Subgroup Performance Analysis. '
        'GRAD model performance stratified by demographic and clinical '
        'subgroups. All metrics computed from ADNI LOOCV predictions. '
        'Age tertiles defined within the ADNI cohort.'
    )

    doc.add_page_break()

    # ── Table S3: Per-Stage Performance ──────────────────────────────
    df = pd.read_csv(f'{TABLES}/supp_table_by_stage.csv')
    df_display = df[['Stage', 'N', 'TP', 'TN', 'FP', 'FN',
                     'Sensitivity', 'Specificity', 'PPV', 'NPV',
                     'Accuracy', 'AUC']].copy()
    df_display.columns = ['Stage', 'N', 'TP', 'TN', 'FP', 'FN',
                          'Sens', 'Spec', 'PPV', 'NPV', 'Acc', 'AUC']
    for c in ['Sens', 'Spec', 'PPV', 'NPV', 'Acc', 'AUC']:
        df_display[c] = df_display[c].apply(lambda x: f'{x:.3f}')
    add_styled_table(
        doc, df_display,
        'Table S3. Per-Stage Performance Metrics. '
        'Performance metrics for each stage of the GRAD algorithm. '
        'Stage 1 (Gatekeeper): univariate p-Tau217 logistic regression. '
        'Stage 2 (Reflex): 6-feature Random Forest classifier applied '
        'to gray zone cases (0.25 \u2264 P \u2264 0.75). '
        'Sens = sensitivity, Spec = specificity, Acc = accuracy.'
    )

    doc.add_page_break()

    # ── Table S4: Operating Points ───────────────────────────────────
    df = pd.read_csv(f'{TABLES}/supp_table_operating_points.csv')
    df['Operating_Point'] = df['Operating_Point'].replace({
        '50% threshold (standard)': '50% Standard',
        '90% Specificity (rule-in)': '90% Spec (Rule-In)',
        '90% Sensitivity (rule-out)': '90% Sens (Rule-Out)',
        'Youden\'s J optimal': 'Youden\'s J',
    })
    df_display = df[['Operating_Point', 'Sensitivity', 'Specificity',
                     'PPV', 'NPV', 'Accuracy', 'LR+', 'LR-']].copy()
    df_display.columns = ['Operating Point', 'Sens', 'Spec',
                          'PPV', 'NPV', 'Acc', 'LR+', 'LR\u2212']
    for c in ['Sens', 'Spec', 'PPV', 'NPV', 'Acc']:
        df_display[c] = df_display[c].apply(lambda x: f'{x:.3f}')
    df_display['LR+'] = df_display['LR+'].apply(lambda x: f'{x:.2f}')
    df_display['LR\u2212'] = df_display['LR\u2212'].apply(lambda x: f'{x:.3f}')
    add_styled_table(
        doc, df_display,
        'Table S4. Clinical Operating Points. '
        'Model performance at clinically relevant operating points. '
        'The 90% sensitivity threshold (rule-out, P < 0.240) and 90% '
        'specificity threshold (rule-in, P > 0.674) define the clinical '
        'decision boundaries. Youden\'s J index identifies the '
        'threshold maximizing sensitivity + specificity \u2212 1.'
    )

    doc.add_page_break()

    # ── Table S5: Threshold Sensitivity ──────────────────────────────
    df = load_threshold_sweep()
    # Columns come from the corrected Gatekeeper sweep, which reports gray-zone
    # and overall AUC instead of the old table's classified-negative/positive
    # counts. Select by name so a change in the source table cannot silently
    # mis-label columns.
    df_display = pd.DataFrame({
        'Low Thr': df['Low_Threshold'].map('{:.2f}'.format),
        'High Thr': df['High_Threshold'].map('{:.2f}'.format),
        'Resolved N': df['resolved_n'].astype(int).astype(str),
        'Resol. Rate': df['Resolution_Rate'].map('{:.1%}'.format),
        'Resol. Acc': df['Resolved_Accuracy'].map('{:.1%}'.format),
        'GZ N': df['Gray_Zone_N'].astype(int).astype(str),
        'GZ AUC': df['Gray_Zone_AUC'].map('{:.3f}'.format),
        'Overall AUC': df['Overall_AUC'].map('{:.3f}'.format),
    })
    add_styled_table(
        doc, df_display,
        'Table S5. Gatekeeper Threshold Sensitivity Analysis. '
        'Each row represents a combination of lower (rule-out) and upper '
        '(rule-in) probability thresholds. Resol. Rate = proportion of '
        'patients classified without Stage 2 or PET. Resol. Acc = accuracy '
        'among resolved cases. GZ = gray zone. The selected operating point '
        'is 0.25/0.75.'
    )

    doc.add_page_break()

    # ── Table S6: Cost Simulation ────────────────────────────────────
    df = pd.read_csv(f'{TABLES}/supp_table_cost_simulation.csv', index_col=0)
    df_display = df[['Strategy', 'Total_Cost', 'Cost_Per_Patient',
                     'PET_Scans', 'Plasma_Tests', 'Plasma_Unit_Cost',
                     'Savings_vs_PET_%']].copy()
    df_display['Total_Cost'] = df_display['Total_Cost'].apply(
        lambda x: f'${x/1e6:.1f}M')
    df_display['Cost_Per_Patient'] = df_display['Cost_Per_Patient'].apply(
        lambda x: f'${x:,.0f}')
    df_display['Plasma_Unit_Cost'] = df_display['Plasma_Unit_Cost'].apply(
        lambda x: f'${x:,.0f}' if x > 0 else '\u2014')
    df_display['PET_Scans'] = df_display['PET_Scans'].apply(
        lambda x: f'{int(x):,}')
    df_display['Plasma_Tests'] = df_display['Plasma_Tests'].apply(
        lambda x: f'{int(x):,}')
    df_display['Savings_vs_PET_%'] = df_display['Savings_vs_PET_%'].apply(
        lambda x: f'{x:.1f}%' if x > 0 else 'Ref')
    df_display.columns = ['Strategy', 'Total', 'Per Capita',
                          'PET Scans', 'Plasma', 'Plasma $', 'Savings']
    add_styled_table(
        doc, df_display,
        'Table S6. Cost Impact Simulation. '
        'Cost impact simulation for a projected 10,000-patient cohort. '
        'Unit costs reflect differentiated plasma pricing: single-analyte '
        'p-Tau217 ($350, CMS 2024 CLFS) versus the full GRAD panel '
        '(p-Tau217, GFAP, A\u03b242/40; $600). PET: $3,000/scan '
        '(CMS 2024 PFS). MRI assumed already obtained for the '
        'Staged + MRI strategy.'
    )

    doc.add_page_break()

    # ── Table S7: NfL Ablation ───────────────────────────────────────
    df = pd.read_csv(f'{TABLES}/supp_table_nfl_ablation.csv')
    df['model'] = df['model'].replace({
        'Base (no NfL)': 'Base (6 features)',
        'Base + NfL_Z': '+ NfL',
        'Base + NfL_Z + nfl_age_interaction': '+ NfL + NfL\u00d7Age',
    })
    df_display = df[['model', 'n_features', 'auc', 'accuracy',
                     'sensitivity', 'specificity', 'brier_score',
                     'delta_auc']].copy()
    df_display.columns = ['Model', 'N', 'AUC', 'Accuracy',
                          'Sensitivity', 'Specificity', 'Brier',
                          '\u0394AUC']
    for c in ['AUC', 'Accuracy', 'Sensitivity', 'Specificity',
              'Brier', '\u0394AUC']:
        df_display[c] = df_display[c].apply(lambda x: f'{x:.3f}')
    add_styled_table(
        doc, df_display,
        'Table S7. NfL Feature Ablation Analysis. '
        'Feature ablation analysis evaluating the contribution of '
        'neurofilament light (NfL) to the Reflex model. Base model '
        'includes the 6 selected features (p-Tau217, GFAP, '
        'Tau\u2013A\u03b242/40 divergence ratio, GFAP \u00d7 p-Tau217 '
        'interaction, age, APOE \u03b54). Adding NfL improved AUC by '
        '0.011, deemed insufficient to justify additional complexity. '
        'All metrics from ADNI LOOCV.'
    )

    doc.save(OUTPUT)
    print(f'Saved: {OUTPUT}')


if __name__ == '__main__':
    main()
