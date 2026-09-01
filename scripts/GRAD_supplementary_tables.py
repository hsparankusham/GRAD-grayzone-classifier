#!/usr/bin/env python3
"""
Supplementary material for the ACTN submission: Tables S1-S9 and Figure S1.
============================================================================

Numbered in order of first mention in the manuscript:

    S1  Analyte availability by cohort and assay platform          (2.2, 2.3)
    S2  Complete-panel sensitivity analysis, no imputation         (2.2)
    S3  Reference-anchoring sensitivity                            (2.3)
    S4  Cost model unit costs and assumptions                      (2.7)
    S5  Diagnostic performance by stage and cohort                 (3.2)
    S6  Operating-point sweep, ADNI                                (3.2)
    S7  Stage 2 associations and model coefficients                (3.4)
    S8  Feature ablation in the A4 + LEARN gray zone               (3.4)
    S9  Predictive values across assumed Ab prevalence             (Discussion)
    Figure S1  Full-pipeline AUC by subgroup, ADNI                 (3.2)

Emits two files from one set of numbers:

    manuscript/SUPPLEMENTARY_TABLES.md          working copy for the repo
    manuscript/GRAD_ACTN_Supplementary.docx     submission copy

Journal formatting for the .docx: tables double spaced in Cambria 12; titles,
notes and figure legends double spaced in Helvetica 12. Wide tables sit in their
own landscape sections - they will not fit portrait at 12 pt.

    python3 scripts/GRAD_supplementary_tables.py
"""

import json

import numpy as np
import pandas as pd
from sklearn.metrics import roc_auc_score, brier_score_loss

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from _grad_paths import RESULTS, PROJECT_ROOT
from run_grad_impaired_logistic import (
    boot_ci, load_adni, load_a4, harmonise, engineer, fit_stage1, fit_stage2,
    score_stage2, IMPAIRED_DX, FEATURES, GK_LOW, GK_HIGH,
)

TABLE_FONT, LEGEND_FONT, SIZE = 'Cambria', 'Helvetica', Pt(12)

PRETTY = {'pTau217_Z': 'p-Tau217', 'tau_ab42_diff': 'tau−Aβ42/40 ratio',
          'GFAP_Z': 'GFAP', 'AGE': 'Age', 'APOE4_carrier': 'APOE ε4 carrier',
          'gfap_tau_interaction': 'GFAP × p-Tau217', 'AB4240_log': 'Aβ42/40',
          'NfL_Z': 'NfL'}


# ------------------------------------------------------------------ helpers ---

def fmt(v, pct=False, dec=3):
    if v is None or (isinstance(v, float) and not np.isfinite(v)):
        return '—'
    return f'{100 * v:.1f}' if pct else f'{v:.{dec}f}'


def minus(s):
    return s.replace('-', '−')


def binary(y, p, thr=.5):
    pred = (p >= thr).astype(int)
    tp = int(((pred == 1) & (y == 1)).sum()); fn = int(((pred == 0) & (y == 1)).sum())
    tn = int(((pred == 0) & (y == 0)).sum()); fp = int(((pred == 1) & (y == 0)).sum())
    sens = tp / (tp + fn) if tp + fn else np.nan
    spec = tn / (tn + fp) if tn + fp else np.nan
    return dict(sens=sens, spec=spec,
                ppv=tp / (tp + fp) if tp + fp else np.nan,
                npv=tn / (tn + fn) if tn + fn else np.nan,
                lr_pos=sens / (1 - spec) if spec < 1 else np.inf,
                lr_neg=(1 - sens) / spec if spec > 0 else np.nan,
                brier=brier_score_loss(y, p))


# ------------------------------------------------------------ table content ---

ASSAY_LABEL = {'UPENN': 'Lumipulse G (Fujirebio, UPENN)',
               'Janssen': 'Janssen immunoassay'}


def table_s1(adni, a4):
    """Analyte availability by cohort and platform, over the analysed cohorts."""
    adni = adni.loc[adni.index[adni['DX'].isin(IMPAIRED_DX)]]
    def counts(df, label):
        return [label, f'{len(df):,}',
                f"{df['pTau217_raw'].notna().sum():,}",
                f"{df['GFAP_raw'].notna().sum():,}",
                f"{df['AB42_40_ratio'].notna().sum():,}",
                f"{df['APOE4_carrier'].notna().sum():,}"]

    rows = [counts(adni[adni['assay'] == a],
                   'ADNI — ' + ASSAY_LABEL.get(a, a))
            for a in sorted(adni['assay'].dropna().unique())]
    rows.append(counts(adni, 'ADNI — all'))
    rows.append(counts(a4, 'A4 + LEARN — Lilly MSD / Roche Elecsys'))
    return rows


def complete_panel(adni, a4):
    """Table S2: external validation restricted to complete-panel participants."""
    idx = adni.index[adni['DX'].isin(IMPAIRED_DX)]
    adni_h = engineer(harmonise(adni, adni))
    a4_h = engineer(harmonise(a4, a4, cn_only=False))
    tr = adni_h.loc[idx]
    y = tr['amyloid_positive'].values
    gk = fit_stage1(tr, y)
    p = gk.predict_proba(tr[['pTau217_Z']].values)[:, 1]
    gzt = (p >= GK_LOW) & (p <= GK_HIGH)
    m2, sc, med, used = fit_stage2(tr[gzt], y[gzt], features=FEATURES)

    p4 = gk.predict_proba(a4_h[['pTau217_Z']].values)[:, 1]
    gz = (p4 >= GK_LOW) & (p4 <= GK_HIGH)
    f = p4.copy()
    f[gz] = score_stage2(m2, sc, med, used, a4_h[gz])
    y4 = a4_h['amyloid_positive'].values
    comp = a4_h[FEATURES].notna().all(axis=1).values

    rows = []
    for label, mask in [('All participants', np.ones(len(y4), bool)),
                        ('Complete panel', comp),
                        ('Imputed ≥1 analyte', ~comp)]:
        ci = boot_ci(y4[mask], f[mask], roc_auc_score)
        g = gz & mask
        rows.append([label, f'{int(mask.sum()):,}', fmt(y4[mask].mean(), True),
                     f'{roc_auc_score(y4[mask], f[mask]):.3f} '
                     f'({ci[0]:.3f}–{ci[1]:.3f})',
                     f'{roc_auc_score(y4[g], f[g]):.3f}',
                     fmt(((f[mask] >= .5).astype(int) == y4[mask]).mean(), True)])

    rng = np.random.RandomState(42)
    ic, ii = np.where(comp)[0], np.where(~comp)[0]
    d = []
    for _ in range(2000):
        bc = rng.choice(ic, len(ic), True)
        bi = rng.choice(ii, len(ii), True)
        if len(set(y4[bc])) < 2 or len(set(y4[bi])) < 2:
            continue
        d.append(roc_auc_score(y4[bc], f[bc]) - roc_auc_score(y4[bi], f[bi]))
    d = np.array(d)
    diff = roc_auc_score(y4[comp], f[comp]) - roc_auc_score(y4[~comp], f[~comp])
    return rows, (diff, np.percentile(d, 2.5), np.percentile(d, 97.5),
                  100 * y4[comp].mean(), 100 * y4[~comp].mean())


def table_s4(cost):
    u = cost['unit_costs']
    return [['Plasma p-Tau217, single analyte', f"${u['plasma']:,.0f}",
             'Every patient, except under universal Aβ-PET'],
            ['Multi-analyte reflex add-on', f"${u['addon']:,.0f}",
             'Only patients routed to Stage 2'],
            ['Aβ-PET scan', f"${u['pet']:,.0f}",
             'Universal strategy: every patient. Staged strategies: '
             'indeterminate cases only'],
            ['Simulated cohort size', f"{u['n']:,}", 'All strategies']]


def table_s6(op):
    rows = []
    for key, label in [('rule_out', 'Rule-out (90% sensitivity target)'),
                       ('rule_in', 'Rule-in (90% specificity target)')]:
        o = op[key]
        rows.append([label, f"{o['threshold']:.3f}", fmt(o['sensitivity'], True),
                     fmt(o['specificity'], True), fmt(o['ppv'], True),
                     fmt(o['npv'], True), f"{o['lr_pos']:.2f}",
                     f"{o['lr_neg']:.2f}", fmt(o['accuracy'], True)])
    return rows


def table_s7(uni, multi, model):
    order = ['APOE4_carrier', 'pTau217_Z', 'AB4240_log', 'tau_ab42_diff',
             'AGE', 'GFAP_Z', 'gfap_tau_interaction']
    u = {x['feature']: x for x in uni}
    m = {x['feature']: x for x in multi} if isinstance(multi, list) else {}
    coef = model['coefficients']
    rows = []
    for feat in order:
        a, b = u[feat], m.get(feat)
        rows.append([PRETTY[feat],
                     f"{a['odds_ratio']:.2f} ({a['lo']:.2f}–{a['hi']:.2f})",
                     f"{a['p']:.3f}".lstrip('0') if a['p'] >= .001 else '<.001',
                     f"{b['odds_ratio']:.2f} ({b['lo']:.2f}–{b['hi']:.2f})"
                     if b else '—',
                     minus(f"{coef[feat]:+.3f}")])
    return rows


def stage_rows(pred, d, cohort, s1key, s2key):
    y, p, g = pred['y'].values, pred['final_prob'].values, pred['gk_prob'].values
    res = (pred['stage'] == 'gatekeeper').values
    gz = ~res
    yr, gr = y[res], g[res]
    call = (gr > .75).astype(int)
    tp = int(((call == 1) & (yr == 1)).sum()); fn = int(((call == 0) & (yr == 1)).sum())
    tn = int(((call == 0) & (yr == 0)).sum()); fp = int(((call == 1) & (yr == 0)).sum())
    sens, spec = tp / (tp + fn), tn / (tn + fp)
    s1 = d[s1key]
    ci1 = s1.get('auc_ci') or list(boot_ci(yr, gr, roc_auc_score))
    rows = [(f'{cohort}, Stage 1', int(res.sum()), s1['auc'], ci1, sens, spec,
             s1['ppv'], s1['npv'], sens / (1 - spec), (1 - sens) / spec, np.nan)]
    s2 = d[s2key]
    mm = s2.get('metrics_at_0.5') or s2.get('metrics_at_0_5')
    rows.append((f'{cohort}, Stage 2', s2['n'], s2['auc_grad'],
                 s2['auc_grad_ci'], mm['sensitivity'], mm['specificity'],
                 mm['ppv'], mm['npv'], mm['lr_pos'], mm['lr_neg'],
                 brier_score_loss(y[gz], p[gz])))
    full = binary(y, p)
    pipe = d['adni_pipeline'] if cohort.startswith('ADNI') else d['a4_pipeline']
    rows.append((f'{cohort}, full pipeline', len(y), pipe['auc'],
                 pipe.get('auc_ci'), full['sens'], full['spec'], full['ppv'],
                 full['npv'], full['lr_pos'], full['lr_neg'], full['brier']))
    return rows


# -------------------------------------------------------------------- notes ---

S1_HEAD = ['Cohort and assay platform', 'n', 'p-Tau217', 'GFAP', 'Aβ42/40',
           'APOE genotype']
S1_NOTE = ('Counts are participants with a non-missing value for each analyte. '
           'ADNI p-Tau217 was measured on the Lumipulse G assay (Fujirebio, '
           'University of Pennsylvania) and on the Janssen assay; the Janssen '
           'subset contributed p-Tau217 only. A4 + LEARN p-Tau217 was measured '
           'on the Lilly Research Laboratories MSD immunoassay, with GFAP and '
           'Aβ42/40 on the Roche Elecsys platform in a subset of the cohort. '
           'ADNI counts refer to the 145 cognitively impaired participants who '
           'formed the development cohort. Because the two cohorts used '
           'different p-Tau217 platforms, all '
           'plasma biomarkers were reference-anchored within cohort and '
           'platform before analysis (Section 2.3).')

S2_HEAD = ['A4 + LEARN subgroup', 'n', 'Aβ+, %', 'Full pipeline AUC (95% CI)',
           'Gray zone AUC', 'Accuracy, %']

S3_LABELS = {'cu_abneg': ('Aβ-negative participants (primary)', 'Yes'),
             'low_tercile': ('Lowest p-Tau217 tercile', 'No'),
             'robust_all': ('Whole cohort, median and IQR', 'No'),
             'mean_all': ('Whole cohort, mean and SD', 'No')}
S3_HEAD = ['A4 + LEARN reference definition', 'Uses Aβ status', 'Stage 1 AUC',
           'Gray zone, %', 'Gray zone AUC', 'Full pipeline AUC', 'Accuracy, %']
S3_NOTE = ('The ADNI-fitted model is held completely fixed; only the mu and '
           'sigma used to reference-anchor the validation cohort change. Stage '
           '1 AUC is computed over the whole cohort and is identical to four '
           'decimal places under every definition, because the Stage 1 '
           'probability is a monotone function of raw p-Tau217 whatever '
           'reference is chosen. The final row draws the reference from a '
           'random half of the Aβ-negative participants, repeated 200 times, '
           'and reports the median with the 2.5th to 97.5th percentile range. '
           'The two whole-cohort definitions deliberately mis-anchor the '
           'scale, since 69.6% of the cohort is Aβ-positive, and are included '
           'to bound the worst case: discrimination is preserved, while the '
           '0.5 probability threshold is displaced and accuracy falls '
           'accordingly. AUC, area under the receiver operating characteristic '
           'curve.')

S4_HEAD = ['Cost component', 'Unit cost', 'Applied to']
S4_NOTE = ('Unit costs were drawn from 2024 U.S. Medicare clinical laboratory '
           'and physician fee schedules (references 43-45). Where a range was '
           'plausible we adopted the higher plasma figure and a conservative '
           'Aβ-PET price so that savings are not overstated. Costs are '
           'undiscounted, expressed in 2024 U.S. dollars, and exclude '
           'clinician time, downstream treatment, and the consequences of '
           'misclassification.')

S5_HEAD = ['Cohort and stage', 'n', 'AUC (95% CI)', 'Sensitivity, %',
           'Specificity, %', 'PPV, %', 'NPV, %', 'LR+', 'LR−', 'Brier']
S5_NOTE = ('Stage 1 rows cover only the participants that stage resolves; its '
           'sensitivity and specificity use the 0.25 and 0.75 routing '
           'thresholds as the decision rule. Stage 2 rows cover the gray zone. '
           'Stage 2 and full-pipeline values use a 0.5 probability threshold. '
           'Confidence intervals are percentile bootstrap over participants '
           '(2,000 resamples). ADNI values are leave-one-out cross-validated; '
           'A4 + LEARN values are from the ADNI-fitted model applied without '
           'refitting. AUC, area under the receiver operating characteristic '
           'curve; LR, likelihood ratio; NPV, negative predictive value; PPV, '
           'positive predictive value.')

S6_HEAD = ['Operating point', 'Probability threshold', 'Sensitivity, %',
           'Specificity, %', 'PPV, %', 'NPV, %', 'LR+', 'LR−', 'Accuracy, %']
S6_NOTE = ('Both operating points are derived in the ADNI development cohort '
           '(n = 145) from leave-one-out cross-validated predictions of the '
           'full pipeline, and are reported to show the range over which the '
           'model can be tuned. They have not been validated prospectively and '
           'should not be adopted as fixed thresholds in other populations '
           'without recalibration. LR, likelihood ratio; NPV, negative '
           'predictive value; PPV, positive predictive value.')

S7_HEAD = ['Feature', 'Univariable OR per SD (95% CI)', 'P',
           'Multivariable OR per SD (95% CI)', 'Penalised coefficient']
S7_NOTE = ('All estimates are from the ADNI gray zone (n = 49). Univariable '
           'odds ratios fit each feature alone and are the values quoted in '
           'Section 3.4. Multivariable odds ratios fit all seven features '
           'together without penalisation and are reported so that the two can '
           'be compared directly; their intervals are wide because the '
           'features are correlated and the sample is small. The final column '
           'gives the L2-penalised coefficients of the deployed Stage 2 model '
           'on standardised features, whose magnitudes are plotted in Figure '
           '2E. The negative coefficient for Aβ42/40 reflects that a lower '
           'ratio indicates greater amyloid burden. The univariable estimates '
           'and the penalised coefficients agree in ranking APOE ε4, p-Tau217 '
           'and Aβ42/40 highest. The unpenalised multivariable estimates are '
           'unstable, with p-Tau217 and the tau−Aβ42/40 ratio taking large '
           'coefficients of opposing sign because the two are collinear by '
           'construction; this instability is why the deployed model is '
           'L2-penalised. OR, odds ratio.')

S8_HEAD = ['Model', 'Gray zone AUC', 'Δ AUC vs full model']
S8_NOTE = ('Each row refits Stage 2 on the ADNI gray zone with the stated '
           'feature set and scores the A4 + LEARN gray zone (n = 630) without '
           'refitting. A negative Δ AUC indicates that removing the feature '
           'degraded external discrimination. NfL was evaluated as an eighth '
           'feature and excluded on this basis.')

S9_HEAD = ['Assumed Aβ prevalence, %', 'PPV, %', 'NPV, %']
S9_NOTE = ('Positive and negative predictive value depend on how common Aβ '
           'positivity is in the population being tested, whereas sensitivity '
           'and specificity do not. Values are derived from the full-pipeline '
           'operating point in A4 + LEARN (76.7% sensitivity, 88.6% '
           'specificity) applied at each assumed prevalence by Bayes theorem. '
           'The 30% and 50% rows approximate a primary care memory complaint '
           'population and a specialist memory clinic respectively; the '
           'observed prevalence was 69.6% in A4 + LEARN and 58.6% in ADNI. '
           'NPV, negative predictive value; PPV, positive predictive value.')

FIGS1_NOTE = ('Area under the receiver operating characteristic curve for the '
              'complete GRAD pipeline within clinical subgroups of the ADNI '
              'development cohort (n = 145), with 95% percentile bootstrap '
              'confidence intervals. The dashed line and shaded band show the '
              'whole-cohort estimate and its interval. Intervals are wide for '
              'the smaller strata, notably AD dementia (n = 28) and APOE ε4 '
              'carriers (n = 61), and should be read as imprecision rather '
              'than instability. The corresponding analysis in the A4 + LEARN '
              'validation cohort is Figure 3A; cognitive status cannot be '
              'stratified there because all participants are cognitively '
              'unimpaired.')


def s2_note(diff):
    d, lo, hi, pc, pi = diff
    return ('Participants are those in the external validation cohort with a '
            'non-missing value for all seven Stage 2 features. The ADNI-fitted '
            'model is unchanged; only the participants scored differ. Panel '
            f'availability was informative: Aβ prevalence was {pc:.1f}% among '
            f'complete-panel participants against {pi:.1f}% among those '
            'requiring imputation, so the two subgroups are not '
            'interchangeable and the difference in AUC between them ('
            + minus(f'{d:+.3f}') + '; 95% CI, ' + minus(f'{lo:+.3f}') + ' to '
            + minus(f'{hi:+.3f}') + ') reflects case mix as well as any effect '
            'of imputation. The interval does not exclude a difference of '
            'clinical relevance and should not be read as demonstrating '
            'equivalence.')


SPEC = [
    ('S1', 'Analyte availability by cohort and assay platform', S1_HEAD,
     S1_NOTE, 'landscape', [2.6, .7, 1.2, 1.1, 1.1, 1.6]),
    ('S2', 'Complete-panel sensitivity analysis in A4 + LEARN', S2_HEAD, None,
     'landscape', [2.2, .8, 1.0, 2.5, 1.4, 1.3]),
    ('S3', 'Sensitivity of external validation to the reference-anchoring '
     'definition', S3_HEAD, S3_NOTE, 'landscape',
     [2.5, 1.0, 1.0, 1.35, 1.6, 1.6, 1.35]),
    ('S4', 'Cost model unit costs and assumptions', S4_HEAD, S4_NOTE,
     'portrait', [2.4, 1.2, 2.9]),
    ('S5', 'Diagnostic performance of GRAD by stage and cohort', S5_HEAD,
     S5_NOTE, 'landscape',
     [1.9, .6, 1.55, 1.0, 1.0, .75, .75, .65, .65, .75]),
    ('S6', 'Operating-point sweep in the ADNI development cohort', S6_HEAD,
     S6_NOTE, 'landscape', [2.3, 1.3, 1.1, 1.1, .8, .8, .7, .7, 1.0]),
    ('S7', 'Stage 2 associations and model coefficients', S7_HEAD, S7_NOTE,
     'landscape', [1.9, 2.3, .8, 2.3, 1.7]),
    ('S8', 'Feature ablation in the A4 + LEARN gray zone', S8_HEAD, S8_NOTE,
     'portrait', [3.0, 1.75, 1.75]),
    ('S9', 'Predictive values across assumed Aβ prevalence', S9_HEAD, S9_NOTE,
     'portrait', [2.5, 2.0, 2.0]),
]


# ------------------------------------------------------------------- build ---

def build():
    d = json.load(open(RESULTS / 'grad_v2_numbers.json'))
    rs = json.load(open(RESULTS / 'reference_sensitivity.json'))
    adni_p = pd.read_csv(RESULTS / 'adni_loocv_predictions_v2.csv')
    a4_p = pd.read_csv(RESULTS / 'a4_predictions_v2.csv')
    adni, a4 = load_adni(), load_a4()

    T = {}
    T['S1'] = table_s1(adni, a4)
    T['S2'], diff = complete_panel(adni, a4)

    s3 = []
    for r in rs['fixed']:
        name, uses = S3_LABELS[r['variant']]
        s3.append([name, uses, f"{r['stage1_auc_all']:.3f}",
                   f"{r['gz_pct']:.1f}", f"{r['stage2_auc']:.3f}",
                   f"{r['pipeline_auc']:.3f}", f"{r['accuracy']:.1f}"])
    h = rs['split_half']

    def rng_(k, dec=3):
        v = h[k]
        return f"{v['median']:.{dec}f} ({v['lo']:.{dec}f}–{v['hi']:.{dec}f})"

    s3.append(['Random half of Aβ-negative participants', 'Yes',
               f"{h['stage1_auc_all']['median']:.3f}", rng_('gz_pct', 1),
               rng_('stage2_auc'), rng_('pipeline_auc'), rng_('accuracy', 1)])
    T['S3'] = s3

    T['S4'] = table_s4(d['cost'])

    raw = (stage_rows(adni_p, d, 'ADNI', 'adni_stage1', 'adni_stage2')
           + stage_rows(a4_p, d, 'A4 + LEARN', 'a4_stage1', 'a4_stage2'))
    T['S5'] = [[n, f'{c:,}',
                f'{a:.3f} ({ci[0]:.3f}–{ci[1]:.3f})' if ci else f'{a:.3f}',
                fmt(se, True), fmt(sp, True), fmt(pp, True), fmt(nv, True),
                fmt(lp, dec=2), fmt(ln, dec=2), fmt(br)]
               for (n, c, a, ci, se, sp, pp, nv, lp, ln, br) in raw]

    T['S6'] = table_s6(d['adni_operating'])
    T['S7'] = table_s7(d['univariate_or'], d.get('odds_ratios'),
                       d['stage2_model'])

    abl = d['a4_ablation']
    base = [a for a in abl if a['variant'] == 'full 7-feature model'][0]
    s8 = [['Full seven-feature model', f"{base['auc']:.3f}", '—']]
    for a in sorted([a for a in abl if a['variant'].startswith('drop')],
                    key=lambda a: a['delta']):
        s8.append([f"Without {PRETTY[a['variant'].replace('drop ', '')]}",
                   f"{a['auc']:.3f}", minus(f"{a['delta']:+.3f}")])
    add = [a for a in abl if a['variant'] == 'add NfL_Z'][0]
    s8.append(['Plus NfL', f"{add['auc']:.3f}", minus(f"{add['delta']:+.3f}")])
    T['S8'] = s8

    T['S9'] = [[f"{r['prevalence'] * 100:.0f}", f"{r['ppv'] * 100:.1f}",
                f"{r['npv'] * 100:.1f}"] for r in d['prevalence_table']]
    return T, diff


# ------------------------------------------------------------------ output ---

def write_md(T, diff):
    out = ['# Supplementary Material', '']
    for key, title, head, note, _, _ in SPEC:
        note = s2_note(diff) if key == 'S2' else note
        out += [f'## Supplementary Table {key}. {title}', '',
                '| ' + ' | '.join(head) + ' |', '|' + '---|' * len(head)]
        out += ['| ' + ' | '.join(r) + ' |' for r in T[key]]
        out += ['', note or '', '', '---', '']
    out += ['## Supplementary Figure S1. Full-pipeline discrimination by '
            'subgroup, ADNI development cohort', '', FIGS1_NOTE, '']
    path = PROJECT_ROOT / 'manuscript' / 'SUPPLEMENTARY_TABLES.md'
    path.write_text('\n'.join(out))
    return path


def style_run(run, font, bold=False):
    run.bold = bold
    run.font.size = SIZE
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), font)


def double(par):
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    pf.space_before = pf.space_after = Pt(0)
    return par


def legend(doc, label, text='', bold_label=True):
    par = double(doc.add_paragraph())
    if label:
        style_run(par.add_run(label), LEGEND_FONT, bold=bold_label)
    if text:
        style_run(par.add_run((' ' if label else '') + text), LEGEND_FONT)
    return par


def hrule(cell, edges):
    b = OxmlElement('w:tcBorders')
    for e in edges:
        el = OxmlElement(f'w:{e}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '8')
        el.set(qn('w:color'), '000000')
        b.append(el)
    cell._element.tcPr.append(b)


def add_table(doc, head, rows, widths):
    t = doc.add_table(rows=1, cols=len(head))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, hh in enumerate(head):
        c = t.rows[0].cells[j]
        style_run(double(c.paragraphs[0]).add_run(hh), TABLE_FONT, bold=True)
        hrule(c, ('top', 'bottom'))
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            style_run(double(cells[j].paragraphs[0]).add_run(v), TABLE_FONT)
            if i == len(rows) - 1:
                hrule(cells[j], ('bottom',))
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Inches(w)
    return t


def section(doc, orient):
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if orient == 'landscape':
        s.orientation = WD_ORIENT.LANDSCAPE
        if s.page_width < s.page_height:
            s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Inches(0.7)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        if s.page_width > s.page_height:
            s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Inches(1.0)
    return s


def write_docx(T, diff):
    doc = Document()
    s = doc.sections[0]
    s.left_margin = s.right_margin = Inches(1.0)
    legend(doc, 'Supplementary Material')
    legend(doc, 'GRAD: a two-stage plasma biomarker algorithm for the p-Tau217 '
                'gray zone', bold_label=False)

    for key, title, head, note, orient, widths in SPEC:
        note = s2_note(diff) if key == 'S2' else note
        section(doc, orient)
        legend(doc, f'Supplementary Table {key}.', title + '.')
        add_table(doc, head, T[key], widths)
        if note:
            legend(doc, '', note)

    section(doc, 'portrait')
    fig = (PROJECT_ROOT / 'results' / 'figures' / 'panels'
           / 'figS1_subgroups_adni.png')
    doc.add_picture(str(fig), width=Inches(6.0))
    doc.paragraphs[-1].alignment = 1
    legend(doc, 'Supplementary Figure S1.',
           'Full-pipeline discrimination by subgroup, ADNI development cohort. '
           + FIGS1_NOTE)

    path = PROJECT_ROOT / 'manuscript' / 'GRAD_ACTN_Supplementary.docx'
    doc.save(path)
    return path


def main():
    T, diff = build()
    print('wrote', write_md(T, diff))
    print('wrote', write_docx(T, diff))
    for key, *_ in SPEC:
        print(f'  {key}: {len(T[key])} rows')


if __name__ == '__main__':
    main()
